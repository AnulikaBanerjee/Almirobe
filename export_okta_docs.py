#!/usr/bin/env python3
"""Generate styled DOCX resume and cover letter for Okta Manager, Engineering application."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Twips, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
import copy

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
SLATE = RGBColor(0x33, 0x3A, 0x45)
MUTED = RGBColor(0x5A, 0x63, 0x72)
ACCENT = RGBColor(0x2C, 0x5F, 0x8A)
RULE = "1B3A5F"

RESUME_PATH = "/Users/anulika/Downloads/Anulika_Banerjee_Okta_Manager_Engineering_Resume.docx"
LETTER_PATH = "/Users/anulika/Downloads/Anulika_Banerjee_Okta_Manager_Engineering_Cover_Letter.docx"
RESUME_WS = "/Users/anulika/Downloads/Wardrobe/Anulika_Banerjee_Okta_Manager_Engineering_Resume.docx"
LETTER_WS = "/Users/anulika/Downloads/Wardrobe/Anulika_Banerjee_Okta_Manager_Engineering_Cover_Letter.docx"


def set_run_font(run, name="Calibri", size=10.5, bold=False, italic=False, color=SLATE):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=4, line=1.08):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_bottom_border(paragraph, color=RULE, size="12"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def set_cell_shading(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="D0D7DE"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def clear_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)


def build_resume():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # Name
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=2, line=1.0)
    r = name.add_run("ANULIKA BANERJEE")
    set_run_font(r, size=20, bold=True, color=NAVY)

    # Headline
    headline = doc.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(headline, before=0, after=1, line=1.05)
    r = headline.add_run("Senior Software Engineer | Tech Lead (Engineering Leadership)")
    set_run_font(r, size=11, bold=True, color=ACCENT)

    # Tagline
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(tag, before=0, after=2, line=1.05)
    r = tag.add_run("Backend  ·  Distributed Systems  ·  Java & Spring Boot  ·  12 Years")
    set_run_font(r, size=10, color=MUTED)

    # Contact
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=8, line=1.05)
    r = contact.add_run(
        "anu.banerjee08@gmail.com  ·  +91 7411288983  ·  linkedin.com/in/anulikabanerjee-b60561b1  ·  Greater Bengaluru"
    )
    set_run_font(r, size=9.5, color=MUTED)
    add_bottom_border(contact, size="18")

    def section_heading(text):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=10, after=4, line=1.0)
        r = p.add_run(text.upper())
        set_run_font(r, size=11, bold=True, color=NAVY)
        add_bottom_border(p, size="10")
        return p

    def body_para(text, after=6):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=after, line=1.08)
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=SLATE)
        return p

    def bullet(text, after=2):
        p = doc.add_paragraph(text, style="List Bullet")
        set_paragraph_spacing(p, before=0, after=after, line=1.05)
        p.paragraph_format.left_indent = Inches(0.2)
        for run in p.runs:
            set_run_font(run, size=10.5, color=SLATE)
        return p

    def job_title(title):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=8, after=0, line=1.0)
        r = p.add_run(title)
        set_run_font(r, size=11, bold=True, color=NAVY)
        return p

    def job_meta(text, italic=False):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=1 if italic else 3, line=1.0)
        r = p.add_run(text)
        if italic:
            set_run_font(r, size=10, italic=True, color=MUTED)
        else:
            set_run_font(r, size=10, bold=True, color=ACCENT)
        return p

    # Summary
    section_heading("Professional Summary")
    body_para(
        "Senior Software Engineer with 12 years designing and delivering enterprise-grade Java/Spring Boot "
        "software at scale. Operates in a Tech Lead capacity: owns end-to-end planning and delivery for "
        "medium-to-large initiatives, mentors engineers through 1:1s, conducts coding interviews to hire talent, "
        "and partners with managers on performance reviews. Hands-on with distributed systems (Kafka, asynchronous "
        "processing), Azure cloud data stores, JWT/OAuth2 authentication patterns, Agile/Scrum, and CI/CD. Strong "
        "cross-functional collaborator with Product Managers, Designers, and Principal Engineers."
    )

    # Key Achievements as styled metric cards (table)
    section_heading("Key Achievements")
    metrics = [
        ("~$3M", "Incremental revenue from Sam's Club TBC initiative"),
        ("80%", "Drop in manual order-flow investigation"),
        ("90%", "Reduction in monitoring effort via alerts"),
        ("~$150K/yr", "Azure cloud savings from rightsizing"),
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, (value, label) in enumerate(metrics):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "F4F7FA")
        set_cell_borders(cell, "C5D0DC")
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p1, before=4, after=0, line=1.0)
        r1 = p1.add_run(value)
        set_run_font(r1, size=12, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p2, before=0, after=4, line=1.0)
        r2 = p2.add_run(label)
        set_run_font(r2, size=8, color=MUTED)

    # Extra achievement as bullet under metrics
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=4, after=0, line=1.0)
    bullet("Built and trained an Accenture engineering team through interviewing and onboarding", after=2)

    # Experience
    section_heading("Professional Experience")

    job_title("Senior Software Engineer | Tech Lead capacity")
    job_meta("Walmart Global Tech India – Sam's Club  |  May 2022 – Present  |  Bengaluru")
    job_meta("Supply Chain & Fulfillment  ·  Distributed Systems  ·  Event-Driven Architecture", italic=True)

    walmart_bullets = [
        "Owned end-to-end delivery for Sam's Garage and Tire & Battery Center (TBC) initiatives—estimations, project planning, and technical oversight—driving on-time feature delivery with Product Managers, Designers, and Principal Engineers.",
        "Architected Java/Spring Boot Kafka microservices handling millions of daily events for order orchestration and fulfillment, applying idempotency, REST contracts, asynchronous processing, and distributed-system design patterns.",
        "Implemented authentication flows using JWT, App2App authentication, and OAuth2 (Gmail integration) within Walmart enterprise services.",
        "Contributed to Sam's Club TBC fulfillment workflows that generated ~$3M in incremental revenue across participating clubs.",
        "Built Elasticsearch indexes (Kibana) over order-flow events to surface stuck, delayed, and failed orders—reducing manual investigation effort by 80%.",
        "Implemented automated alerts on order-flow health, Kafka lag, and error-rate thresholds across Splunk, Prometheus, Grafana, and OpenObserve—cutting reactive monitoring effort by 90%.",
        "Extracted Azure Cost Management and Advisor reports to rightsize compute and storage—delivering ~$150K in annual cloud savings.",
        "Owned production RCA on multi-region Kubernetes with Oracle DB, Azure Cosmos DB, and Azure Blob Storage; shipped via JUnit, Maven, Git, and enterprise CI/CD.",
        "Mentored 4–6 engineers through coaching and 1:1s; conducted coding interviews to hire engineers; partnered with Manager on performance reviews; facilitated Scrum.",
    ]
    for b in walmart_bullets:
        bullet(b)

    job_title("Application Development Team Lead")
    job_meta("Accenture  |  Dec 2021 – May 2022")
    for b in [
        "Built the team by conducting interviews, hiring engineers, and training them for delivery readiness.",
        "Led a team of 5 engineers delivering enterprise integrations connecting EBS, ERP, Oracle DB, FTP, and UCM with robust error handling and audit trails.",
        "Owned technical quality, error-handling patterns, accurate estimations, and stakeholder alignment during UAT and production releases.",
    ]:
        bullet(b)

    job_title("Application Development Senior Analyst")
    job_meta("Accenture  |  Apr 2020 – Dec 2021")
    bullet(
        "Developed and maintained Core Java and Oracle Integration Cloud (OIC) flows; contributed to architecture reviews and team onboarding."
    )

    job_title("Senior Associate Consultant / Associate Consultant")
    job_meta("Infosys  |  Feb 2018 – Mar 2020")
    bullet(
        "Designed Java and OSB web services (SOAP/REST); optimized integration flows; gathered client requirements; mentored junior developers."
    )

    job_title("Software Engineer")
    job_meta("Cognizant Technology Solutions  |  Jun 2014 – Feb 2018")
    bullet(
        "Built enterprise integration pipelines using Java, REST/SOAP services, and Oracle DB/JMS adapters; recognized with Shining Star (x2), The Incredibles, and Insta awards."
    )

    # Skills
    section_heading("Technical Skills")
    skills = [
        ("Languages", "Java 8+ (primary), SQL"),
        ("Backend", "Spring Boot, Spring Framework, REST APIs, Microservices, Event-Driven Architecture, CQRS, Kafka, Design Patterns, Caching, Asynchronous Processing, Multi-threading"),
        ("Authentication & Security", "JWT, OAuth2 (including Gmail integration), App2App authentication"),
        ("Data", "Oracle DB, Azure Cosmos DB, Azure Blob Storage, Elasticsearch, Kibana"),
        ("Cloud & Ops", "Azure, Kubernetes, Maven, Git, JUnit, Walmart Enterprise CI/CD, Prometheus, Grafana, OpenObserve, Splunk, Distributed Tracing"),
        ("Leadership & Delivery", "Tech Lead capacity, Hiring (coding interviews), Mentoring & 1:1s, Performance review input, Estimation & project planning, Scrum, Cross-functional collaboration (Product / Design / Principal Engineers), Incident RCA"),
    ]
    for label, value in skills:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=2, line=1.05)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=SLATE)

    # Education
    section_heading("Education & Certifications")
    body_para(
        "B.Tech, Computer Science & Engineering – F.I.E.M. Kolkata  |  2010 – 2014",
        after=2,
    )
    for cert in [
        "Oracle Cloud Platform Application Integration 2020 Certified Specialist",
        "Oracle Cloud Infrastructure Foundations 2020 Certified Associate",
        "CutShort Certified: Java – Basic",
    ]:
        bullet(cert, after=1)

    doc.save(RESUME_PATH)
    doc.save(RESUME_WS)
    return RESUME_PATH


def build_cover_letter():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Header block
    name = doc.add_paragraph()
    set_paragraph_spacing(name, before=0, after=2, line=1.0)
    r = name.add_run("ANULIKA BANERJEE")
    set_run_font(r, size=16, bold=True, color=NAVY)

    contact = doc.add_paragraph()
    set_paragraph_spacing(contact, before=0, after=2, line=1.05)
    r = contact.add_run("anu.banerjee08@gmail.com  ·  +91 7411288983  ·  Greater Bengaluru")
    set_run_font(r, size=10, color=MUTED)

    linkedin = doc.add_paragraph()
    set_paragraph_spacing(linkedin, before=0, after=10, line=1.05)
    r = linkedin.add_run("linkedin.com/in/anulikabanerjee-b60561b1")
    set_run_font(r, size=10, color=MUTED)
    add_bottom_border(linkedin, size="14")

    subject = doc.add_paragraph()
    set_paragraph_spacing(subject, before=8, after=10, line=1.1)
    r = subject.add_run("Re: Manager, Engineering – Authentication Experience")
    set_run_font(r, size=11, bold=True, color=NAVY)

    greeting = doc.add_paragraph()
    set_paragraph_spacing(greeting, before=0, after=8, line=1.15)
    r = greeting.add_run("Dear Hiring Manager,")
    set_run_font(r, size=11, color=SLATE)

    paragraphs = [
        "I am applying for the Manager, Engineering role on Okta’s Authentication Experience team. I am drawn to work that balances robust security with a seamless user experience—and to leading engineers who deliver that balance with urgency and craftsmanship.",
        "I bring 12 years of software development experience, including hands-on design and delivery of enterprise Java/Spring Boot systems at scale. At Walmart Global Tech (Sam’s Club), I work as a Senior Software Engineer in a Tech Lead capacity. I have owned end-to-end delivery for medium-to-large initiatives such as Sam’s Garage and the Tire & Battery Center (TBC) program—driving estimations, planning, technical oversight, and on-time feature delivery while partnering closely with Product Managers, Designers, and Principal Engineers on product and technical decisions.",
        "Technically, I design and develop Java/Spring Boot Kafka microservices that process millions of events daily, with strong grounding in object-oriented design, distributed systems concepts (asynchronous processing, reliability patterns), debugging complex production Java systems, and clean, maintainable code backed by unit tests and CI/CD. I operate critical paths on Azure (Cosmos DB, Blob Storage), Oracle, and Kubernetes. At Walmart, I implemented authentication flows using JWT, App2App authentication, and OAuth2 (Gmail integration). I care deeply about quality, operational reliability, and a high bar for secure, maintainable software.",
        "As a leader, I mentor engineers through coaching and 1:1s, conduct coding interviews to hire talent, and partner with my Manager on performance reviews. At Accenture, as Application Development Team Lead, I built a team by interviewing, hiring, and training engineers, and led a team of five through enterprise integration delivery with clear quality standards and stakeholder alignment. That mix of still being hands-on in enterprise Java while owning planning, people coaching, and cross-functional delivery is how I lead.",
        "I would welcome the opportunity to bring that ownership to Authentication Experience—aligning technical strategy with Product, Architecture, QA, Documentation, and UX; managing major feature lifecycles from ideation to delivery; and mentoring engineers who shape how Workforce and Customer users authenticate safely and seamlessly.",
        "Thank you for your consideration.",
    ]

    for text in paragraphs:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=0, after=8, line=1.15)
        r = p.add_run(text)
        set_run_font(r, size=11, color=SLATE)

    close = doc.add_paragraph()
    set_paragraph_spacing(close, before=6, after=2, line=1.15)
    r = close.add_run("Sincerely,")
    set_run_font(r, size=11, color=SLATE)

    sign = doc.add_paragraph()
    set_paragraph_spacing(sign, before=12, after=0, line=1.15)
    r = sign.add_run("Anulika Banerjee")
    set_run_font(r, size=11, bold=True, color=NAVY)

    doc.save(LETTER_PATH)
    doc.save(LETTER_WS)
    return LETTER_PATH


if __name__ == "__main__":
    r = build_resume()
    c = build_cover_letter()
    print(f"Resume: {r}")
    print(f"Cover letter: {c}")
